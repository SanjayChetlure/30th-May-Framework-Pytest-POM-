import openpyxl



workbook=openpyxl.load_workbook("D:\Python\Workspace\8thNov_pytestFramework\TestData\SwagLab.xlsx")
sheet=workbook['Sheet2']

#Apr1
data=sheet["A1"].value
print(data)
print(sheet["D1"].value)
print(sheet["A2"].value)

print("---")

#Apr2
data1=sheet.cell(1,1).value
print(data1)

print(sheet.cell(1,4).value)

#convert data to int
data2=sheet.cell(2,1).value
data2_int=int(data2)
